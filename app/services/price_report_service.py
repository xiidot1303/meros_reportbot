"""Build the price-list xlsx files that the bot hands to clients.

The price list is identical for every client — `ProductPrice` carries no client
FK — so the files are generated once per sync (~2s each for ~20k rows) and every
client is served the same file. Generating per click would rebuild an identical
workbook and, because openpyxl is CPU-bound, would block the bot's event loop
for everyone else while it ran.

Files are written atomically: a temp file in the same directory followed by
`os.replace`, so a click landing mid-write is served the previous complete file
rather than a truncated one.
"""

import logging
import os
import tempfile
from datetime import datetime

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.models import ProductPrice


logger = logging.getLogger(__name__)


PRICE_LIST_DIR = "files/price_lists"

# Whose price list this is — printed at the top of every sheet.
COMPANY_NAME = "MerosPharm MCHJ"

# Rows the info block occupies above the table (company, type, date, blank).
HEADER_ROWS = 4

# Columns of the generated sheet, in order: (header, model field, width).
COLUMNS = [
    ("Наименование товара", "product_name", 55),
    ("Производитель", "manufacturer", 32),
    ("Кол-во в коробке", "box_quant", 18),
    ("Код карточки", "card_code", 16),
    ("Цена", "price", 16),
    ("Остаток", "quant", 14),
    ("Срок годности", "expiry_date", 16),
]

FIELDS = [field for _, field, _ in COLUMNS]

# Rows are streamed from the DB in chunks so the whole price list is never
# materialised as model instances.
DB_CHUNK_SIZE = 2000


def price_list_filename(price_type_id, generated_at):
    """Timestamped name, so each sync produces a distinct file on disk."""
    return f"price_list_{price_type_id}_{generated_at:%Y%m%d_%H%M%S}.xlsx"


def client_facing_filename(price_type_id, generated_at):
    """What the client sees in Telegram — dated, no internal ids."""
    label = {
        ProductPrice.PREPAYMENT_25: "25_predoplata",
        ProductPrice.PREPAYMENT_100: "100_predoplata",
    }.get(price_type_id, str(price_type_id))
    return f"Прайс-лист_{label}_{generated_at:%d.%m.%Y}.xlsx"


def _format_value(field, value):
    if value is None:
        return None
    if field == "expiry_date":
        return value.strftime("%d.%m.%Y")
    if field in ("box_quant", "price", "quant"):
        return float(value)
    return value


def _write_info_header(sheet, price_type_id, generated_at):
    """Company, price type and data age, above the table.

    Written as label/value pairs in columns A/B so the values stay readable and
    the block survives sorting or filtering the table below it.
    """
    price_type_label = dict(ProductPrice.PRICE_TYPE_CHOICES).get(
        price_type_id, str(price_type_id))

    sheet["A1"] = COMPANY_NAME
    sheet["A1"].font = Font(bold=True, size=14)

    sheet["A2"] = "Тип цены:"
    sheet["B2"] = price_type_label

    sheet["A3"] = "Актуально на:"
    sheet["B3"] = generated_at.strftime("%d.%m.%Y %H:%M")

    for row in (2, 3):
        sheet[f"A{row}"].font = Font(bold=True)
        sheet[f"B{row}"].alignment = Alignment(horizontal="left")


def build_price_list_file(price_type_id, generated_at=None):
    """Write one price type's xlsx and return (path, row_count).

    Returns `(None, 0)` when the price type has no rows — a failed sync should
    not replace a good file with an empty one.
    """
    generated_at = generated_at or timezone.now()

    rows = (
        ProductPrice.objects
        .filter(price_type_id=price_type_id)
        .order_by("product_name", "card_code")
        .values_list(*FIELDS)
    )

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Прайс-лист"

    _write_info_header(sheet, price_type_id, generated_at)

    # HEADER_ROWS leaves a blank spacer row, so the table starts right after it
    header_row = HEADER_ROWS + 1
    for index, (header, _, _) in enumerate(COLUMNS, start=1):
        cell = sheet.cell(row=header_row, column=index, value=header)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.fill = PatternFill("solid", start_color="DDEBF7")

    count = 0
    for row in rows.iterator(chunk_size=DB_CHUNK_SIZE):
        sheet.append([
            _format_value(field, value) for field, value in zip(FIELDS, row)
        ])
        count += 1

    if not count:
        logger.warning(
            "Price type %s has no rows; skipping file generation", price_type_id)
        return None, 0

    # freeze the info block and the column headers, so scrolling keeps both
    sheet.freeze_panes = f"A{header_row + 1}"
    for index, (_, _, width) in enumerate(COLUMNS, start=1):
        sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = width

    os.makedirs(PRICE_LIST_DIR, exist_ok=True)
    path = os.path.join(PRICE_LIST_DIR, price_list_filename(price_type_id, generated_at))

    # Write to a temp file in the same directory, then rename onto the final
    # name: `os.replace` is atomic within a filesystem, so a reader sees either
    # the old file or the new one, never a half-written workbook.
    handle, temp_path = tempfile.mkstemp(dir=PRICE_LIST_DIR, suffix=".xlsx.tmp")
    os.close(handle)
    try:
        workbook.save(temp_path)
        # mkstemp creates the file 0600; the bot may run as another user, so
        # widen it to the usual read permissions before publishing.
        os.chmod(temp_path, 0o644)
        os.replace(temp_path, path)
    except Exception:
        if os.path.exists(temp_path):
            os.remove(temp_path)
        raise

    logger.info("Built price list %s (%s rows)", path, count)
    return path, count


def cleanup_old_price_lists(keep_paths):
    """Delete every previously generated price list except the current ones."""
    if not os.path.isdir(PRICE_LIST_DIR):
        return 0

    keep = {os.path.abspath(path) for path in keep_paths if path}
    removed = 0
    for name in os.listdir(PRICE_LIST_DIR):
        if not (name.startswith("price_list_") or name.endswith(".xlsx.tmp")):
            continue
        path = os.path.abspath(os.path.join(PRICE_LIST_DIR, name))
        if path in keep:
            continue
        try:
            os.remove(path)
            removed += 1
        except OSError:
            logger.warning("Could not delete old price list %s", path, exc_info=True)

    if removed:
        logger.info("Removed %s outdated price list file(s)", removed)
    return removed


def generate_price_list_files():
    """Rebuild both price lists, publish them, and drop the previous files.

    The published record (path, generation time, row count and the cached
    Telegram file_id) lives in `PriceListFile`, which is what the bot reads.
    """
    from app.models import PriceListFile

    generated_at = timezone.now()
    current_paths = []

    for price_type_id in (ProductPrice.PREPAYMENT_25, ProductPrice.PREPAYMENT_100):
        path, count = build_price_list_file(price_type_id, generated_at)
        if not path:
            # Keep the previous file published rather than serving nothing.
            existing = PriceListFile.objects.filter(price_type_id=price_type_id).first()
            if existing and existing.exists:
                current_paths.append(existing.path)
            continue

        PriceListFile.publish(
            price_type_id=price_type_id,
            path=path,
            generated_at=generated_at,
            row_count=count,
        )
        current_paths.append(path)

    cleanup_old_price_lists(current_paths)
    return current_paths
